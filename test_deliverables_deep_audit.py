import os
import openpyxl
import numpy as np
import pandas as pd
from pathlib import Path

ROOT = Path(".").resolve()
TPL_DIR = ROOT / "data" / "raw" / "Annex5_Templates"
RES_DIR = ROOT / "results"

class Colors:
    GREEN = "\033[92m"
    RED = "\033[91m"
    CYAN = "\033[96m"
    RESET = "\033[0m"

def log_pass(msg): print(f"  [{Colors.GREEN}PASS{Colors.RESET}] {msg}")
def log_fail(msg): print(f"  [{Colors.RED}FAIL{Colors.RESET}] {msg}")

def assert_file_topology(filename, expected_sheets):
    f_res = RES_DIR / filename
    f_tpl = TPL_DIR / filename
    assert f_res.exists(), f"交付文件缺失: {f_res}"
    
    wb_r = openpyxl.load_workbook(f_res, data_only=True)
    assert wb_r.sheetnames == expected_sheets, (
        f"{filename} Sheet 结构错误! 期望: {expected_sheets}, 实际: {wb_r.sheetnames}"
    )
    if f_tpl.exists():
        wb_t = openpyxl.load_workbook(f_tpl, data_only=True)
        assert wb_r.sheetnames == wb_t.sheetnames, f"{filename} Sheet 名与官方母版不一致"
    log_pass(f"{filename} 工作表结构 100% 吻合: {wb_r.sheetnames}")
    return f_res

def check_plan_sheet(f_path, sheet_name="计划购电量", expected_rows=334):
    df = pd.read_excel(f_path, sheet_name=sheet_name)
    assert len(df) == expected_rows, f"{sheet_name} 行数错误: 期望 {expected_rows}, 实际 {len(df)}"
    data_mat = df.iloc[:, 1:145].values.astype(float)
    assert not np.isnan(data_mat).any(), f"{sheet_name} 存在 NaN 缺失值"
    assert (data_mat >= -1e-5).all(), f"{sheet_name} 存在负数购电 (禁止向外网倒送)"
    assert data_mat.max() < 3000.0, (
        f"{sheet_name} 单步数值异常偏大 ({data_mat.max():.2f})，疑似误写成功率 kW 而非电量 kWh!"
    )
    log_pass(f"{sheet_name} 时序维度 ({expected_rows}×144) 及电量量纲通过")

def check_battery_sheet(f_path, expected_rows=2004):
    df = pd.read_excel(f_path, sheet_name="充放电量")
    assert len(df) == expected_rows, f"充放电量表行数错误: 期望 {expected_rows}, 实际 {len(df)}"
    
    # 日期显示规范：每日首行展示，其余 5 行留空
    date_col = df.iloc[:, 0]
    for d in range(expected_rows // 6):
        assert pd.notna(date_col.iloc[d * 6]), f"第 {d+1} 天首行日期丢失"
        assert date_col.iloc[d * 6 + 1 : (d + 1) * 6].isna().all(), (
            f"第 {d+1} 天后续 5 行出现冗余日期填充"
        )
    log_pass("充放电量表日期排版合规 (每日仅首行展示)")

    # 初始储电量物理安全与继承核验
    feb1_soc0 = float(df.iloc[0, 5])
    assert 1200.0 <= feb1_soc0 <= 10800.0, (
        f"2月1日 0:00 初始电量越界 ({feb1_soc0:.2f} kWh)! 必须处于 [1200, 10800] 物理安全区间内!"
    )
    log_pass(f"自然暖机初值物理有效性通过 (Feb 1 0:00 = {feb1_soc0:.2f} kWh)")

    # 跨日物理连续性：Day d 24:00 电量 == Day d+1 0:00 电量
    day_ends = df.iloc[1::6, 5].values.astype(float)
    day_starts = df.iloc[6::6, 5].values.astype(float)
    discrepancy = np.abs(day_ends[:-1] - day_starts)
    max_jump = float(np.max(discrepancy))
    assert max_jump < 1e-3, f"出现跨日物理电量跳跃! 最大跳跃 = {max_jump:.4f} kWh"
    log_pass(f"334 天跨日电量无缝衔接断言通过 (最大跳跃误差 = {max_jump:.6f} kWh)")

def check_emergency_sheet(f_path, min_events=50, max_vol=150000.0):
    df = pd.read_excel(f_path, sheet_name="紧急购电量")
    assert list(df.columns) == ["日期", "购电时间段", "购电量"], (
        f"紧急购电量表头错误! 必须严格为 3 列，当前为: {list(df.columns)}"
    )
    assert not df.isnull().any().any(), "紧急购电量存在 NaN 空值或空白占位行"
    assert len(df) >= min_events, f"缺电事件行数过少 ({len(df)})"
    total_vol = float(df["购电量"].sum())
    assert 1000.0 < total_vol < max_vol, f"紧急购电总量异常: {total_vol:.2f} kWh (上限 {max_vol:.2f})"
    log_pass(f"紧急购电事件表规范通过: {len(df)} 笔真实缺电记录，总量 {total_vol:.2f} kWh")

def main():
    print("\n" + "="*80)
    print("      CUMCM 2026 交付成果终审：官方母版对拍与物理运筹硬核断言")
    print("="*80)
    try:
        print(f"\n{Colors.CYAN}>>> [Q1 确定性排产断言]{Colors.RESET}")
        f1 = assert_file_topology("result1.xlsx", ["计划购电量", "充放电量"])
        df1_p = pd.read_excel(f1, sheet_name="计划购电量")
        assert len(df1_p) == 144, "Q1 计划购电量行数非 144"
        log_pass("Q1 单日 144 步长计划排产合规")

        print(f"\n{Colors.CYAN}>>> [Q2 两阶段鲁棒排产断言]{Colors.RESET}")
        f2 = assert_file_topology("result2.xlsx", ["计划购电量", "充放电量", "紧急购电量"])
        check_plan_sheet(f2, "计划购电量")
        check_battery_sheet(f2)
        check_emergency_sheet(f2, min_events=50, max_vol=120000.0)

        print(f"\n{Colors.CYAN}>>> [Q3 闭环滚动 MPC 断言]{Colors.RESET}")
        f3 = assert_file_topology("result3.xlsx", ["计划购电量", "调整购电量", "充放电量", "紧急购电量"])
        check_plan_sheet(f3, "计划购电量")
        check_plan_sheet(f3, "调整购电量")
        check_battery_sheet(f3)
        check_emergency_sheet(f3, min_events=50, max_vol=120000.0)

        print(f"\n{Colors.CYAN}>>> [Q4 动态电价交付断言]{Colors.RESET}")
        f4_2 = assert_file_topology("result4-2.xlsx", ["计划购电量", "充放电量", "紧急购电量"])
        check_plan_sheet(f4_2, "计划购电量")
        check_battery_sheet(f4_2)
        check_emergency_sheet(f4_2, min_events=50, max_vol=130000.0)

        f4_3 = assert_file_topology("result4-3.xlsx", ["计划购电量", "调整购电量", "充放电量", "紧急购电量"])
        check_plan_sheet(f4_3, "计划购电量")
        check_plan_sheet(f4_3, "调整购电量")
        check_battery_sheet(f4_3)
        check_emergency_sheet(f4_3, min_events=50, max_vol=120000.0)

        print("\n" + "="*80)
        print(f"{Colors.GREEN}【终审全部通过 ALL ASSERTIONS PASSED】全部 Excel 交付物均达最高合规水准！{Colors.RESET}")
        print("="*80 + "\n")
    except AssertionError as err:
        print("\n" + "="*80)
        print(f"{Colors.RED}【断言阻断 ASSERTION FAILED】: {err}{Colors.RESET}")
        print("="*80 + "\n")

if __name__ == "__main__":
    main()