import os
import numpy as np
import pandas as pd
from openpyxl import load_workbook

from src.config import DELTA_T, E_INIT

TEMPLATES_DIR = os.path.join(os.path.dirname(__file__), "..", "data", "raw", "Annex5_Templates")
RESULTS_DIR = os.path.join(os.path.dirname(__file__), "..", "results")

# 4-hour aggregation windows: (start_step, end_step) exclusive end.
FOUR_HOUR_WINDOWS = [
    (0, 24), (24, 48), (48, 72), (72, 96), (96, 120), (120, 144),
]


def _ensure_dir(path):
    os.makedirs(path, exist_ok=True)


def export_q1_result(res):
    """
    Populates the blank Annex 5 result1 template and saves to results/result1.xlsx.
    Never overwrites data/raw templates.
    """
    template_path = os.path.join(TEMPLATES_DIR, "result1.xlsx")
    out_path = os.path.join(RESULTS_DIR, "result1.xlsx")
    _ensure_dir(RESULTS_DIR)

    wb = load_workbook(template_path)

    ws1 = wb["计划购电量"]
    for t in range(len(res["p_grid_kw"])):
        ws1.cell(row=2 + t, column=2, value=float(round(res["p_grid_kw"][t] * DELTA_T, 4)))

    ws2 = wb["充放电量"]
    for i, (a, b) in enumerate(FOUR_HOUR_WINDOWS):
        chg = float(np.sum(res["p_chg_kw"][a:b]) * DELTA_T)
        dis = float(np.sum(res["p_dis_kw"][a:b]) * DELTA_T)
        ws2.cell(row=2 + i, column=2, value=float(round(chg, 4)))
        ws2.cell(row=2 + i, column=3, value=float(round(dis, 4)))

    for r in range(1, ws2.max_row + 1):
        for c in range(1, ws2.max_column + 1):
            val = str(ws2.cell(row=r, column=c).value).strip()
            if val in ["0:00", "0:00:00"]:
                ws2.cell(row=r, column=c + 1, value=6000.0)
            elif val in ["24:00", "24:00:00", "0:00+1"]:
                ws2.cell(row=r, column=c + 1, value=6000.0)

    wb.save(out_path)
    print(f"[SUCCESS] Q1 deliverables correctly mapped to: {out_path}")
    return out_path


def export_result(file_name, data_dict):
    """
    Writes generated data arrays to results/ (data/raw/ is read-only).
    If the target file is locked (open in Excel), falls back to a copy
    with a `_generated` suffix instead of crashing.
    """
    _ensure_dir(RESULTS_DIR)
    target_path = os.path.join(RESULTS_DIR, file_name)
    df = pd.DataFrame(data_dict)
    try:
        df.to_excel(target_path, index=False)
        print(f"[SUCCESS] Formatted energy outputs saved to: {target_path}")
    except PermissionError:
        base, ext = os.path.splitext(file_name)
        fallback = os.path.join(RESULTS_DIR, f"{base}_generated{ext}")
        df.to_excel(fallback, index=False)
        print(f"[WARNING] {file_name} is open in Excel; results saved to: {fallback}")